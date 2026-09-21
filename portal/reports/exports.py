import csv
import io
from datetime import datetime

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font

from inventory.models import Dispatch
from reports.period import parse_period_bounds
from sales.models import SyncLog
from accounts.roles import filter_by_accessible_branches


def _make_response(filename, content_type):
    resp = HttpResponse(content_type=content_type)
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def _write_csv(rows, headers):
    resp = _make_response(
        f"report_{datetime.now():%Y%m%d_%H%M%S}.csv", "text/csv"
    )
    writer = csv.writer(resp)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return resp


def _write_xlsx(rows, headers, sheet_title="Report"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    resp = _make_response(
        f"report_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    wb.save(resp)
    return resp


def _export(request, headers, rows, sheet_title="Report"):
    fmt = request.GET.get("format", "csv")
    if fmt == "xlsx":
        return _write_xlsx(rows, headers, sheet_title)
    return _write_csv(rows, headers)


# ── Stock balance ──────────────────────────────────────────────

@login_required
def export_stock_balance(request):
    from reports.views import stock_balance_data

    # Respects the same ?branch= / period filters as the on-screen report.
    _bounds, lines, by_sku, by_branch, totals = stock_balance_data(request)
    view = (request.GET.get("view") or "detail").lower()

    def _tons_cell(value):
        return float(value) if value is not None else ""

    if view == "sku":
        headers = ["Product", "Qty", "Metric Tons", "Stock Value"]
        rows = [
            [r["label"], r["quantity"], _tons_cell(r["tons"]), float(r["stock_value"])]
            for r in by_sku
        ]
        rows.append(
            [
                "Grand Total",
                totals["quantity"],
                _tons_cell(totals["tons"]),
                float(totals["stock_value"]),
            ]
        )
        title = "Stock Balance by Product"
    elif view == "branch":
        headers = ["Branch", "Qty", "Metric Tons", "Stock Value"]
        rows = [
            [r["label"], r["quantity"], _tons_cell(r["tons"]), float(r["stock_value"])]
            for r in by_branch
        ]
        rows.append(
            [
                "Grand Total",
                totals["quantity"],
                _tons_cell(totals["tons"]),
                float(totals["stock_value"]),
            ]
        )
        title = "Stock Balance by Branch"
    else:
        headers = [
            "Customer",
            "Branch",
            "Code",
            "Product",
            "Qty",
            "Metric Tons",
            "Unit Price",
            "Stock Value",
        ]
        rows = [
            [
                r["customer_name"],
                r["branch_name"],
                r["product_code"],
                r["product_name"],
                r["quantity"],
                _tons_cell(r["tons"]),
                float(r["unit_price"]),
                float(r["stock_value"]),
            ]
            for r in lines
        ]
        rows.append(
            [
                "Grand Total",
                "",
                "",
                "",
                totals["quantity"],
                _tons_cell(totals["tons"]),
                "",
                float(totals["stock_value"]),
            ]
        )
        title = "Stock Balance"
    return _export(request, headers, rows, title)

# ── Dispatch report ────────────────────────────────────────────

@login_required
def export_dispatches(request):
    bounds = parse_period_bounds(request)
    dispatches = filter_by_accessible_branches(
        Dispatch.objects.select_related("customer", "branch")
        .prefetch_related("items__product")
        .filter(created_at__gte=bounds["start_dt"], created_at__lt=bounds["end_dt"]),
        request.user,
    )
    status = request.GET.get("status")
    if status:
        dispatches = dispatches.filter(status=status)
    headers = ["Reference", "Customer", "Branch", "Status", "Items"]
    rows = [
        [
            d.reference,
            d.customer.name,
            d.branch.name,
            d.get_status_display(),
            ", ".join(f"{i.product.code}×{i.quantity}" for i in d.items.all()),
        ]
        for d in dispatches
    ]
    return _export(request, headers, rows, "Dispatches")


# ── Sales report (VastAfrica template) ─────────────────────────

def _sales_raw_row(r):
    d = r["date"]
    date_str = f"{d.day}/{d.month}/{d.year}" if d else ""
    return [
        date_str,
        r["item_description"],
        r["group"],
        r["sale_reporting_group"],
        r["account"],
        r["customer_name"],
        float(r["tons"]) if r["tons"] is not None else "",
        r["quantity"],
        float(r["amount"]),
        float(r["unit_price"]),
        r["branch"],
        float(r["psize"]) if r["psize"] is not None else "",
    ]


def _write_sales_xlsx(bounds, raw_rows, by_category, by_sku, by_branch, totals):
    from reports.views import COMPANY_NAME

    wb = Workbook()

    ws_cat = wb.active
    ws_cat.title = "Sales-By-Category"
    ws_cat.append(["Group", "(All)"])
    ws_cat.append([])
    ws_cat.append(["Product Category", "Volume"])
    for cell in ws_cat[3]:
        cell.font = Font(bold=True)
    for r in by_category:
        ws_cat.append([r["label"], float(r["volume"])])
    ws_cat.append(["Grand Total", float(totals["tons"])])
    ws_cat["A" + str(ws_cat.max_row)].font = Font(bold=True)

    ws_sku = wb.create_sheet("Sales-by-SKU")
    ws_sku.append(["PSize", "(All)"])
    ws_sku.append([])
    ws_sku.append(["Product Name", "Volume"])
    for cell in ws_sku[3]:
        cell.font = Font(bold=True)
    for r in by_sku:
        ws_sku.append([r["label"], float(r["volume"])])
    ws_sku.append(["Grand Total", float(totals["tons"])])
    ws_sku["A" + str(ws_sku.max_row)].font = Font(bold=True)

    ws_br = wb.create_sheet("Sales-by-Branch")
    ws_br.append(["Date", "(All)"])
    ws_br.append([])
    ws_br.append(["Branch Sales", "Volume"])
    for cell in ws_br[3]:
        cell.font = Font(bold=True)
    for r in by_branch:
        ws_br.append([r["label"], float(r["volume"])])
    ws_br.append(["Grand Total", float(totals["tons"])])
    ws_br["A" + str(ws_br.max_row)].font = Font(bold=True)

    ws_raw = wb.create_sheet("Raw-Data")
    ws_raw["C1"] = "Inventory Transactions"
    ws_raw["C1"].font = Font(bold=True)
    ws_raw["C2"] = COMPANY_NAME
    ws_raw["A3"] = "Date From :"
    ws_raw["B4"] = f"{bounds['date_from'].day}/{bounds['date_from'].month}/{bounds['date_from'].year}"
    ws_raw["A5"] = "Date To :"
    ws_raw["B5"] = f"{bounds['date_to'].day}/{bounds['date_to'].month}/{bounds['date_to'].year}"
    ws_raw["A6"] = "Inventory Transactions"
    headers = [
        "Date",
        "Item Description",
        "Group",
        "Sale Reporting Group",
        "Account",
        "Customer/ Supplier Name",
        "Tons",
        "Quantity",
        "Amount",
        "Unit Price",
        "Branch",
        "PSize",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws_raw.cell(row=7, column=col, value=header)
        cell.font = Font(bold=True)
    for row_idx, r in enumerate(raw_rows, start=8):
        for col_idx, value in enumerate(_sales_raw_row(r), start=1):
            ws_raw.cell(row=row_idx, column=col_idx, value=value)

    resp = _make_response(
        f"sales_report_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    wb.save(resp)
    return resp


@login_required
def export_sales(request):
    from reports.views import sales_report_data

    bounds, raw_rows, by_category, by_sku, by_branch, totals = sales_report_data(
        request
    )
    fmt = request.GET.get("format", "csv")
    if fmt == "xlsx":
        return _write_sales_xlsx(
            bounds, raw_rows, by_category, by_sku, by_branch, totals
        )

    headers = [
        "Date",
        "Item Description",
        "Group",
        "Sale Reporting Group",
        "Account",
        "Customer/ Supplier Name",
        "Tons",
        "Quantity",
        "Amount",
        "Unit Price",
        "Branch",
        "PSize",
    ]
    rows = [_sales_raw_row(r) for r in raw_rows]
    return _write_csv(rows, headers)


# ── Customer stock summary ─────────────────────────────────────

@login_required
def export_customer_stock(request):
    from reports.views import customer_stock_summary_data

    data = customer_stock_summary_data(request.user)
    headers = [
        "Customer",
        "Branches",
        "Units",
        "Metric Tons",
        "Stock Value",
        "Tons Sold",
    ]
    rows = [
        [
            r["customer_name"],
            r["branch_count"],
            r["total_qty"],
            float(r["tons"]),
            float(r["stock_value"]),
            float(r["total_tons_sold"]),
        ]
        for r in data
    ]
    return _export(request, headers, rows, "Customer Stock")


# ── Sync report ────────────────────────────────────────────────

@login_required
def export_sync(request):
    bounds = parse_period_bounds(request)
    logs = filter_by_accessible_branches(
        SyncLog.objects.select_related("branch").filter(
            created_at__gte=bounds["start_dt"], created_at__lt=bounds["end_dt"]
        ),
        request.user,
    )[:100]
    headers = ["When", "Branch", "Status", "Sale ID", "Message"]
    rows = [
        [
            str(l.created_at),
            l.branch.name,
            l.get_status_display(),
            l.external_sale_id,
            l.message,
        ]
        for l in logs
    ]
    return _export(request, headers, rows, "Sync Report")


# ── Stock movement ─────────────────────────────────────────────

@login_required
def export_stock_movement(request):
    from reports.views import stock_movement_data

    data = stock_movement_data(request)
    group_by = data["group_by"]
    if group_by == "branch":
        headers = ["Customer", "Branch", "Code", "Product", "Opening", "Dispatched", "Sold", "Shrinkage", "Closing"]
    else:
        headers = ["Customer", "Code", "Product", "Opening", "Dispatched", "Sold", "Shrinkage", "Closing"]

    export_rows = []
    for r in data["rows"]:
        row = [r["customer"].name]
        if group_by == "branch":
            row.append(r["branch"].name)
        row.extend(
            [
                r["product"].code,
                r["product"].name,
                r["opening"],
                r["dispatched"],
                r["sold"],
                r["shrinkage"],
                r["closing"],
            ]
        )
        export_rows.append(row)

    return _export(request, headers, export_rows, "Stock Movement")


# ── Dispatch warehouse ─────────────────────────────────────────

@login_required
def export_dispatch_warehouse(request):
    from reports.views import dispatch_warehouse_data

    _bounds, groups, _grand_total = dispatch_warehouse_data(request)
    headers = ["Category", "Code", "Product", "Customer", "Branch", "Qty dispatched"]
    export_rows = []
    for group in groups:
        for r in group["rows"]:
            export_rows.append(
                [
                    group["category"],
                    r["code"],
                    r["name"],
                    r["customer"],
                    r["branch"],
                    r["quantity"],
                ]
            )
        export_rows.append(
            [group["category"], "", "Category total", "", "", group["total_qty"]]
        )
    return _export(request, headers, export_rows, "Dispatch Warehouse")


# ── Customer tonnage (POS GRV) ─────────────────────────────────

@login_required
def export_customer_tonnage(request):
    from reports.views import customer_tonnage_data

    _bounds, data, totals, error = customer_tonnage_data(request)

    headers = [
        "Customer",
        "Purchased Units",
        "Purchased Tons",
        "Sold Tons",
        "Variance Tons",
        "Amount",
        "GRV Count",
        "Line Count",
    ]
    rows = [
        [
            r["customer_name"],
            r["quantity"],
            float(r["purchased_tons"]),
            float(r["sold_tons"]),
            float(r["variance_tons"]),
            float(r["amount"]),
            r["grv_count"],
            r["line_count"],
        ]
        for r in data
    ]
    if not rows and error:
        return _export(request, ["Error"], [[error]], "Customer Tonnage")

    rows.append(
        [
            "Grand Total",
            totals["quantity"],
            float(totals["purchased_tons"]),
            float(totals["sold_tons"]),
            float(totals["variance_tons"]),
            float(totals["amount"]),
            "",
            "",
        ]
    )
    return _export(request, headers, rows, "Customer Tonnage")
